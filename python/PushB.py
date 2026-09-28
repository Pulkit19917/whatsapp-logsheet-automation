"""
push_tab_b.py

Pushes every row in MASTER_LOG_SHEET_B out to its Source File Name,
writing into the correct block/row/column.

Configuration:
    config/config.json

Main configuration:
    paths.master_file
    paths.folder_b

    push_tab_b.dry_run
    push_tab_b.master_sheet
    push_tab_b.master_header_row
    push_tab_b.max_header_search_rows
    push_tab_b.max_date_search_rows
    push_tab_b.max_block_width
    push_tab_b.max_left_scan
    push_tab_b.max_empty_streak

HOW IT LOCATES THE RIGHT CELL:

  1. Reads Source File Name, UID, Block Start Column, Date, Shift,
     and reading columns from MASTER_LOG_SHEET_B.

  2. Opens the Source File Name in the B folder.

  3. Automatically selects the correct month sheet from the row Date.

  4. Searches down the Block Start Column for "Date".

  5. Uses that row as the block header row.

  6. Matches standardized field names to target columns.

  7. Finds the Date + Shift row.

  8. Writes the values into the correct cells.

SAFETY:

    dry_run = true

should be used first to review the complete plan.

Set:

    dry_run = false

only when you are ready to actually write data.
"""

import io
import os
import re
import time
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.cell.cell import MergedCell


# ============================================================
# COLORAMA
# ============================================================

try:
    from colorama import init as colorama_init, Fore, Style

    colorama_init()

except ImportError:

    class _NoColor:

        def __getattr__(self, name):
            return ""

    Fore = Style = _NoColor()


# ============================================================
# PROJECT / CONFIGURATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

CONFIG_FILE = PROJECT_ROOT / "config" / "config.json"


def load_config():
    """Load the central project configuration."""

    if not CONFIG_FILE.exists():

        print()
        print("=" * 70)
        print("ERROR: config.json was not found.")
        print()
        print("Expected:")
        print(CONFIG_FILE)
        print("=" * 70)

        raise SystemExit(1)

    try:

        with open(
            CONFIG_FILE,
            "r",
            encoding="utf-8"
        ) as f:

            return __import__("json").load(f)

    except Exception as e:

        print()
        print("=" * 70)
        print("ERROR: Could not read config.json")
        print()
        print(e)
        print("=" * 70)

        raise SystemExit(1)


CONFIG = load_config()

PATH_CONFIG = CONFIG.get(
    "paths",
    {}
)

PUSH_CONFIG = CONFIG.get(
    "push_tab_b",
    {}
)


# ============================================================
# CONFIG HELPERS
# ============================================================

def resolve_config_path(
    value,
    default_relative
):
    """
    Resolve a configured path.

    Absolute paths are used directly.

    Relative paths are resolved from PROJECT_ROOT.
    """

    if value is None or str(value).strip() == "":
        value = default_relative

    path = Path(
        str(value)
    )

    if path.is_absolute():
        return path

    return PROJECT_ROOT / path


def config_bool(
    value,
    default=False
):
    """Safely convert config value to bool."""

    if isinstance(value, bool):
        return value

    if value is None:
        return default

    if isinstance(value, str):

        value = value.strip().lower()

        if value in {
            "1",
            "true",
            "yes",
            "y",
            "on"
        }:
            return True

        if value in {
            "0",
            "false",
            "no",
            "n",
            "off"
        }:
            return False

    return default


def config_int(
    value,
    default
):
    """Safely convert config value to integer."""

    try:
        return int(value)

    except (
        TypeError,
        ValueError
    ):
        return default


# ============================================================
# PATHS
# ============================================================

MASTER_FILE = resolve_config_path(
    PATH_CONFIG.get(
        "master_file"
    ),
    "data/Master_Sheet2.xlsm"
)

FOLDER_B = resolve_config_path(
    PATH_CONFIG.get(
        "folder_b"
    ),
    "data/Logsheet-All Site2"
)


# ============================================================
# PUSH TAB B SETTINGS
# ============================================================

DRY_RUN = config_bool(
    PUSH_CONFIG.get(
        "dry_run"
    ),
    True
)

MASTER_SHEET_B = PUSH_CONFIG.get(
    "master_sheet",
    "MASTER_LOG_SHEET_B"
)

MASTER_HEADER_ROW = config_int(
    PUSH_CONFIG.get(
        "master_header_row"
    ),
    1
)

MAX_HEADER_SEARCH_ROWS = config_int(
    PUSH_CONFIG.get(
        "max_header_search_rows"
    ),
    25
)

MAX_DATE_SEARCH_ROWS = config_int(
    PUSH_CONFIG.get(
        "max_date_search_rows"
    ),
    70
)

MAX_BLOCK_WIDTH = config_int(
    PUSH_CONFIG.get(
        "max_block_width"
    ),
    20
)

MAX_LEFT_SCAN = config_int(
    PUSH_CONFIG.get(
        "max_left_scan"
    ),
    6
)

MAX_EMPTY_STREAK = config_int(
    PUSH_CONFIG.get(
        "max_empty_streak"
    ),
    300
)


# ============================================================
# TIMER
# ============================================================

_T0 = time.time()


def log(msg):
    """Timestamped progress line."""

    print(
        f"{Fore.CYAN}"
        f"[{time.time() - _T0:6.1f}s]"
        f"{Style.RESET_ALL} "
        f"{msg}",
        flush=True
    )


# ============================================================
# FILE READING
# ============================================================

def read_file_bytes(path):
    """
    Read the entire file sequentially.

    This is faster for network shares than allowing openpyxl
    to perform many small network reads.
    """

    t = time.time()

    with open(path, "rb") as f:
        data = f.read()

    log(
        f"  read {os.path.basename(path)} "
        f"({len(data) / 1e6:.1f} MB in "
        f"{time.time() - t:.1f}s)"
    )

    return data


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize(text):

    if text is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(text).strip().lower()
    )


BLANK_MARKERS = {
    "-",
    "--",
    "—",
    "–",
    "n/a",
    "na"
}


def is_blank_marker(value):

    if value is None:
        return True

    return normalize(value) in BLANK_MARKERS


# ============================================================
# SAVE WITH RETRY
# ============================================================

def save_with_retry(
    wb,
    filepath
):
    """
    Save workbook.

    If Excel has the file open, wait for user to close it
    and retry.
    """

    while True:

        try:

            wb.save(filepath)

            return

        except PermissionError:

            print(
                f"\n{Fore.YELLOW}"
                f"Can't save - "
                f"'{os.path.basename(filepath)}' "
                f"looks like it's open in Excel."
                f"{Style.RESET_ALL}"
            )

            input(
                "  Close it, then press Enter "
                "to try saving again..."
            )


# ============================================================
# DATE NORMALIZATION
# ============================================================

def normalize_date(value):

    if value is None:
        return None

    if isinstance(
        value,
        datetime
    ):
        return value.date()

    if (
        hasattr(value, "year")
        and hasattr(value, "month")
    ):
        return value

    text = str(value).strip()

    formats = [
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%d/%m/%y",
        "%d-%m-%y"
    ]

    for fmt in formats:

        try:

            return datetime.strptime(
                text,
                fmt
            ).date()

        except ValueError:

            continue

    return None


# ============================================================
# FIND MONTH SHEET
# ============================================================

def find_matching_sheet(
    wb,
    target_date
):

    if len(wb.sheetnames) == 1:

        return (
            wb.sheetnames[0],
            None
        )

    d = normalize_date(
        target_date
    )

    if d is None:

        return (
            None,
            f"Could not parse date "
            f"'{target_date}'"
        )

    candidates = [
        d.strftime("%b-%y"),
        d.strftime("%B-%y"),
        d.strftime("%b-%Y"),
        d.strftime("%B-%Y"),
        d.strftime("%b%y"),
        d.strftime("%b%y") + "-Update"
    ]

    normalized_sheets = {
        normalize(s): s
        for s in wb.sheetnames
    }

    for candidate in candidates:

        key = normalize(
            candidate
        )

        if key in normalized_sheets:

            return (
                normalized_sheets[key],
                None
            )

    return (
        None,
        f"Workbook has sheets "
        f"{wb.sheetnames} but none matched "
        f"patterns {candidates}."
    )


# ============================================================
# FIND BLOCK HEADER ROW
# ============================================================

def find_block_header_row(
    ws,
    start_col_letter
):

    try:

        col_idx = column_index_from_string(
            str(start_col_letter)
        )

    except Exception:

        return None

    for r in range(
        1,
        MAX_HEADER_SEARCH_ROWS + 1
    ):

        value = ws.cell(
            row=r,
            column=col_idx
        ).value

        if normalize(value) == "date":

            return r

    return None


# ============================================================
# FIND BLOCK HEADER MAP
# ============================================================

def find_block_header_map(
    ws,
    start_col_letter,
    header_row
):
    """
    Build:

        {
            normalized header:
            column number
        }

    Scans both directions from Date.

    Stops at the first blank column.
    """

    start_idx = column_index_from_string(
        str(start_col_letter)
    )

    header_map = {}

    # --------------------------------------------------------
    # RIGHT OF DATE
    # --------------------------------------------------------

    for offset in range(
        MAX_BLOCK_WIDTH
    ):

        col = start_idx + offset

        value = ws.cell(
            row=header_row,
            column=col
        ).value

        if (
            value is None
            or str(value).strip() == ""
        ):

            break

        key = normalize(
            value
        )

        if key:

            header_map.setdefault(
                key,
                col
            )

    # --------------------------------------------------------
    # LEFT OF DATE
    # --------------------------------------------------------

    for offset in range(
        1,
        MAX_LEFT_SCAN + 1
    ):

        col = start_idx - offset

        if col < 1:
            break

        value = ws.cell(
            row=header_row,
            column=col
        ).value

        if (
            value is None
            or str(value).strip() == ""
        ):

            break

        key = normalize(
            value
        )

        if key:

            header_map.setdefault(
                key,
                col
            )

    return header_map


# ============================================================
# FIND DATE + SHIFT ROW
# ============================================================

def find_date_shift_row(
    ws,
    date_col_idx,
    shift_col_idx,
    header_row,
    target_date,
    target_shift
):

    target_d = normalize_date(
        target_date
    )

    seen = []

    for r in range(
        header_row + 1,
        header_row
        + 1
        + MAX_DATE_SEARCH_ROWS
    ):

        cell_date_raw = ws.cell(
            row=r,
            column=date_col_idx
        ).value

        cell_shift = (
            ws.cell(
                row=r,
                column=shift_col_idx
            ).value
            if shift_col_idx
            else None
        )

        cell_d = normalize_date(
            cell_date_raw
        )

        if cell_d is not None:

            seen.append(
                (
                    r,
                    cell_date_raw,
                    cell_shift
                )
            )

        shift_ok = (
            shift_col_idx is None
            or normalize(cell_shift)
            == normalize(target_shift)
        )

        if (
            cell_d is not None
            and target_d is not None
            and cell_d == target_d
            and shift_ok
        ):

            return r, seen

    return None, seen


# ============================================================
# BUILD FILE INDEX
# ============================================================

def build_file_index(
    root_folder
):
    """
    Scan the B folder once.

    Returns:

        {
            "filename":
                [full paths]
        }
    """

    t = time.time()

    root = os.path.abspath(
        str(root_folder)
    )

    index = {}

    count = 0

    if not os.path.isdir(root):

        return index

    for current_root, _dirs, files in os.walk(
        root
    ):

        for file_name in files:

            lower_name = file_name.lower()

            if (
                not lower_name.endswith(
                    (
                        ".xlsx",
                        ".xlsm"
                    )
                )
                or lower_name.startswith("~$")
            ):

                continue

            stem = os.path.splitext(
                lower_name
            )[0]

            full_path = os.path.join(
                current_root,
                file_name
            )

            index.setdefault(
                stem,
                []
            ).append(
                full_path
            )

            count += 1

    for paths in index.values():

        paths.sort(
            key=lambda x: (
                x.lower().count(
                    os.sep
                ),
                x.lower()
            )
        )

    log(
        f"  indexed {count} workbook(s) "
        f"in {time.time() - t:.1f}s"
    )

    return index


_warned_duplicates = set()


# ============================================================
# FIND SOURCE FILE
# ============================================================

def find_source_file_recursive(
    root_folder,
    source_name,
    index
):
    """
    Find Source File Name anywhere under FOLDER_B.

    Supports:

        filename.xlsx
        filename.xlsm
        filename
        relative path
    """

    if not source_name:

        return None

    root = os.path.abspath(
        str(root_folder)
    )

    raw = str(
        source_name
    ).strip().strip('"')

    raw = raw.replace(
        "/",
        os.sep
    )

    # --------------------------------------------------------
    # Direct relative path
    # --------------------------------------------------------

    direct_candidates = [
        os.path.join(
            root,
            raw
        )
    ]

    if not raw.lower().endswith(
        (
            ".xlsx",
            ".xlsm"
        )
    ):

        direct_candidates.extend(
            [
                os.path.join(
                    root,
                    raw + ".xlsx"
                ),
                os.path.join(
                    root,
                    raw + ".xlsm"
                )
            ]
        )

    for candidate in direct_candidates:

        if os.path.isfile(candidate):

            return os.path.abspath(
                candidate
            )

    # --------------------------------------------------------
    # Filename matching
    # --------------------------------------------------------

    wanted_stem = os.path.splitext(
        os.path.basename(
            raw
        ).lower()
    )[0]

    matches = index.get(
        wanted_stem,
        []
    )

    if not matches:

        return None

    if (
        len(matches) > 1
        and wanted_stem
        not in _warned_duplicates
    ):

        _warned_duplicates.add(
            wanted_stem
        )

        print(
            f"{Fore.YELLOW}"
            f"WARNING: multiple files found "
            f"for '{source_name}'."
            f"{Style.RESET_ALL}"
        )

        print(
            f"  Using: {matches[0]}"
        )

        for duplicate in matches[1:]:

            print(
                f"  Other match: {duplicate}"
            )

    return os.path.abspath(
        matches[0]
    )


# ============================================================
# READ MASTER
# ============================================================

def read_master_rows():
    """
    Read MASTER_LOG_SHEET_B in read-only mode.

    Returns:

        rows,
        pushed_column
    """

    log(
        f"Opening master (read-only): "
        f"{MASTER_FILE.name}"
    )

    wb = load_workbook(
        io.BytesIO(
            read_file_bytes(
                MASTER_FILE
            )
        ),
        data_only=True,
        read_only=True
    )

    if MASTER_SHEET_B not in wb.sheetnames:

        print(
            f"\nERROR: sheet "
            f"'{MASTER_SHEET_B}' "
            f"not found in {MASTER_FILE}"
        )

        print(
            "Actual sheet names:",
            wb.sheetnames
        )

        raise SystemExit(1)

    log(
        f"Reading {MASTER_SHEET_B} ..."
    )

    ws = wb[
        MASTER_SHEET_B
    ]

    row_iter = ws.iter_rows(
        min_row=MASTER_HEADER_ROW,
        values_only=True
    )

    header_vals = next(
        row_iter,
        None
    ) or ()

    headers = {}

    for c, v in enumerate(
        header_vals,
        start=1
    ):

        if v:

            headers[
                normalize(v)
            ] = c

    # --------------------------------------------------------
    # Header synonyms
    # --------------------------------------------------------

    COLUMN_SYNONYMS = {

        "block start column": [
            "block start column",
            "block_start",
            "block start",
            "blockstart"
        ]
    }

    def col(name):

        key = normalize(
            name
        )

        variants = COLUMN_SYNONYMS.get(
            key,
            [key]
        )

        for variant in variants:

            if variant in headers:

                return headers[
                    variant
                ]

        return None

    # --------------------------------------------------------
    # Required columns
    # --------------------------------------------------------

    required = [
        "date",
        "source file name",
        "uid",
        "block start column"
    ]

    missing = [
        h
        for h in required
        if col(h) is None
    ]

    if missing:

        print()
        print(
            f"ERROR: {MASTER_SHEET_B} "
            f"is missing expected column(s):"
        )

        print(
            missing
        )

        print()
        print(
            f"Columns found on row "
            f"{MASTER_HEADER_ROW}:"
        )

        print(
            list(headers.keys())
        )

        raise SystemExit(1)

    # --------------------------------------------------------
    # Pushed column
    # --------------------------------------------------------

    pushed_col = col(
        "Pushed"
    )

    # --------------------------------------------------------
    # Fields
    # --------------------------------------------------------

    field_names = [

        "Starting HMR",

        "Closing HMR",

        "Starting KMR",

        "Closing KMR",

        "Breakdown Start Time",

        "Breakdown End Time",

        "Breakdown Hours",

        "Diesel Received (Ltr)",

        "Diesel Issued (Ltr)",

        "HMR at Diesel Fill",

        "Hours Run Since Last Fill",

        "Avg. Diesel Consumption (Ltr/Hr)",

        "Container Start Count",

        "Container End Count"
    ]

    field_cols = {
        fname: col(fname)
        for fname in field_names
        if col(fname)
    }

    # --------------------------------------------------------
    # Master columns
    # --------------------------------------------------------

    date_c = col("date")
    uid_c = col("uid")

    src_c = col(
        "source file name"
    )

    blk_c = col(
        "block start column"
    )

    shift_c = col(
        "shift"
    )

    def get(row, c):

        if (
            c
            and c - 1 < len(row)
        ):

            return row[c - 1]

        return None

    # --------------------------------------------------------
    # Read rows
    # --------------------------------------------------------

    rows = []

    r = MASTER_HEADER_ROW

    empty_streak = 0

    for row in row_iter:

        r += 1

        date_val = get(
            row,
            date_c
        )

        uid_val = get(
            row,
            uid_c
        )

        # ----------------------------------------------------
        # Empty row
        # ----------------------------------------------------

        if (
            not date_val
            and not uid_val
        ):

            empty_streak += 1

            if (
                empty_streak
                >= MAX_EMPTY_STREAK
            ):

                break

            continue

        empty_streak = 0

        # ----------------------------------------------------
        # Incomplete row
        # ----------------------------------------------------

        if (
            not date_val
            or not uid_val
        ):

            continue

        # ----------------------------------------------------
        # Already pushed
        # ----------------------------------------------------

        if (
            pushed_col
            and get(
                row,
                pushed_col
            ) == "Y"
        ):

            continue

        # ----------------------------------------------------
        # Build entry
        # ----------------------------------------------------

        entry = {

            "row": r,

            "date": date_val,

            "uid": uid_val,

            "source_file": get(
                row,
                src_c
            ),

            "block_start_col": get(
                row,
                blk_c
            ),

            "shift": (
                get(
                    row,
                    shift_c
                )
                if shift_c
                else None
            )
        }

        # ----------------------------------------------------
        # Read fields
        # ----------------------------------------------------

        for fname, c in field_cols.items():

            value = get(
                row,
                c
            )

            if not is_blank_marker(
                value
            ):

                entry[
                    fname
                ] = value

        rows.append(
            entry
        )

    wb.close()

    log(
        f"  read {len(rows)} "
        f"unpushed row(s) "
        f"(scanned down to sheet row {r})"
    )

    return (
        rows,
        pushed_col
    )


# ============================================================
# DATE INPUT
# ============================================================

def parse_dd_mm(
    text,
    year
):
    """
    Parse DD/MM against supplied year.
    """

    text = text.strip().replace(
        "-",
        "/"
    )

    try:

        return datetime.strptime(
            text,
            "%d/%m"
        ).replace(
            year=year
        ).date()

    except ValueError:

        return None


def prompt_for_date(
    rows
):
    """
    Ask user which date to process.
    """

    available_dates = sorted(
        {
            normalize_date(
                r["date"]
            )
            for r in rows
            if normalize_date(
                r["date"]
            )
        }
    )

    if available_dates:

        years = [
            d.year
            for d in available_dates
        ]

        default_year = max(
            set(years),
            key=years.count
        )

    else:

        default_year = datetime.now().year

    # --------------------------------------------------------
    # Display dates
    # --------------------------------------------------------

    if available_dates:

        print(
            "Dates currently in "
            "the sheet (unpushed rows):"
        )

        for d in available_dates:

            count = sum(
                1
                for r in rows
                if normalize_date(
                    r["date"]
                ) == d
            )

            print(
                f"  - {d.strftime('%d/%m')} "
                f"({count} row(s))"
            )

    else:

        print(
            f"{Fore.RED}"
            f"No valid dates found "
            f"in the sheet."
            f"{Style.RESET_ALL}"
        )

    # --------------------------------------------------------
    # Ask date
    # --------------------------------------------------------

    while True:

        text = input(
            f"\nEnter the date to update "
            f"(DD/MM, assumes "
            f"{default_year}): "
        ).strip()

        parsed = parse_dd_mm(
            text,
            default_year
        )

        if parsed is None:

            print(
                f"{Fore.RED}"
                f"Couldn't understand "
                f"'{text}' as a date."
                f"{Style.RESET_ALL}"
            )

            continue

        if (
            available_dates
            and parsed
            not in available_dates
        ):

            print(
                f"{Fore.YELLOW}"
                f"Warning: "
                f"{parsed.strftime('%d/%m')} "
                f"isn't in the list above."
                f"{Style.RESET_ALL}"
            )

            confirm = input(
                "Proceed anyway? (y/n): "
            ).strip().lower()

            if confirm != "y":

                continue

        return parsed


# ============================================================
# MAIN
# ============================================================

def main(
    remembered_date=None
):

    print(
        "=" * 64
    )

    print(
        f"  push_tab_b.py - "
        f"{'DRY RUN (no files changed)' if DRY_RUN else 'LIVE (files WILL be written)'}"
    )

    print(
        "=" * 64
    )

    print()
    print(
        "Project root:"
    )
    print(
        f"  {PROJECT_ROOT}"
    )

    print()
    print(
        "Config:"
    )
    print(
        f"  {CONFIG_FILE}"
    )

    print()
    print(
        "Master:"
    )
    print(
        f"  {MASTER_FILE}"
    )

    print()
    print(
        "B Source Folder:"
    )
    print(
        f"  {FOLDER_B}"
    )

    # --------------------------------------------------------
    # Validate master
    # --------------------------------------------------------

    if not MASTER_FILE.exists():

        print()
        print(
            f"{Fore.RED}"
            f"ERROR: Master file not found:"
            f"{Style.RESET_ALL}"
        )

        print(
            MASTER_FILE
        )

        return remembered_date

    # --------------------------------------------------------
    # Validate B folder
    # --------------------------------------------------------

    if not FOLDER_B.exists():

        print()
        print(
            f"{Fore.RED}"
            f"ERROR: B source folder not found:"
            f"{Style.RESET_ALL}"
        )

        print(
            FOLDER_B
        )

        return remembered_date

    if not FOLDER_B.is_dir():

        print()
        print(
            f"{Fore.RED}"
            f"ERROR: B source path is not "
            f"a folder:"
            f"{Style.RESET_ALL}"
        )

        print(
            FOLDER_B
        )

        return remembered_date

    # --------------------------------------------------------
    # Read master
    # --------------------------------------------------------

    rows, pushed_col = read_master_rows()

    print()

    print(
        f"{MASTER_SHEET_B} has "
        f"{len(rows)} unpushed row(s) "
        f"across all dates."
    )

    # --------------------------------------------------------
    # Date
    # --------------------------------------------------------

    if remembered_date is not None:

        target_date = remembered_date

        print()

        print(
            "Re-running for same date:"
        )

        print(
            target_date.strftime(
                "%d/%m"
            )
        )

    else:

        target_date = prompt_for_date(
            rows
        )

    # --------------------------------------------------------
    # Filter by date
    # --------------------------------------------------------

    rows = [
        r
        for r in rows
        if normalize_date(
            r["date"]
        ) == target_date
    ]

    if not rows:

        print()

        print(
            f"{Fore.RED}"
            f"No rows found for "
            f"{target_date.strftime('%d/%m')}."
            f"{Style.RESET_ALL}"
        )

        return target_date

    print()

    print(
        f"{MASTER_SHEET_B}: "
        f"{len(rows)} row(s) "
        f"to process for "
        f"{target_date.strftime('%d/%m')}"
    )

    # --------------------------------------------------------
    # Index B folder
    # --------------------------------------------------------

    log(
        "Indexing the B logsheet folder "
        "(once) ..."
    )

    file_index = build_file_index(
        FOLDER_B
    )

    # --------------------------------------------------------
    # Master write workbook
    # --------------------------------------------------------

    master_wb_write = None
    src_ws_write = None

    if (
        not DRY_RUN
        and pushed_col
    ):

        log(
            "Opening master for writing "
            "(sets Pushed flags)..."
        )

        master_wb_write = load_workbook(
            io.BytesIO(
                read_file_bytes(
                    MASTER_FILE
                )
            ),
            keep_vba=str(
                MASTER_FILE
            ).lower().endswith(
                ".xlsm"
            )
        )

        src_ws_write = (
            master_wb_write[
                MASTER_SHEET_B
            ]
        )

    # --------------------------------------------------------
    # Target workbooks
    # --------------------------------------------------------

    open_targets = {}

    open_errors = {}

    results = []

    # --------------------------------------------------------
    # Process
    # --------------------------------------------------------

    log(
        f"Processing {len(rows)} row(s) ..."
    )

    for entry in rows:

        uid = entry[
            "uid"
        ]

        filename = (
            entry["source_file"]
            or "(none)"
        )

        result = {

            "uid": uid,

            "file": filename,

            "status": None,

            "reason": None,

            "fields_written": 0,

            "fields_skipped": 0
        }

        # ----------------------------------------------------
        # Required information
        # ----------------------------------------------------

        if (
            not entry["source_file"]
            or not entry["block_start_col"]
        ):

            result["status"] = "SKIP"

            result["reason"] = (
                "missing Source File Name "
                "or Block Start Column"
            )

            results.append(
                result
            )

            continue

        # ----------------------------------------------------
        # Find source workbook
        # ----------------------------------------------------

        filepath = find_source_file_recursive(
            FOLDER_B,
            filename,
            file_index
        )

        if filepath is None:

            result["status"] = "SKIP"

            result["reason"] = (
                f"file not found anywhere "
                f"under '{FOLDER_B}'"
            )

            results.append(
                result
            )

            continue

        # ----------------------------------------------------
        # Open workbook once
        # ----------------------------------------------------

        if filepath not in open_targets:

            log(
                f"Opening "
                f"{os.path.basename(filepath)} ..."
            )

            try:

                open_targets[
                    filepath
                ] = load_workbook(
                    io.BytesIO(
                        read_file_bytes(
                            filepath
                        )
                    )
                )

            except Exception as e:

                open_targets[
                    filepath
                ] = None

                open_errors[
                    filepath
                ] = (
                    f"{type(e).__name__}: "
                    f"{e}"
                )

        target_wb = open_targets[
            filepath
        ]

        if target_wb is None:

            result["status"] = "SKIP"

            result["reason"] = (
                "could not open file "
                f"({open_errors[filepath]})"
            )

            results.append(
                result
            )

            continue

        # ----------------------------------------------------
        # Find month sheet
        # ----------------------------------------------------

        sheet_found, err = find_matching_sheet(
            target_wb,
            entry["date"]
        )

        if sheet_found is None:

            result["status"] = "SKIP"

            result["reason"] = err

            results.append(
                result
            )

            continue

        target_ws = target_wb[
            sheet_found
        ]

        # ----------------------------------------------------
        # Find block header
        # ----------------------------------------------------

        block_start = entry[
            "block_start_col"
        ]

        header_row = find_block_header_row(
            target_ws,
            block_start
        )

        if header_row is None:

            result["status"] = "SKIP"

            result["reason"] = (
                f"no 'Date' header in "
                f"column {block_start} "
                f"of '{sheet_found}'"
            )

            results.append(
                result
            )

            continue

        # ----------------------------------------------------
        # Build header map
        # ----------------------------------------------------

        header_map = find_block_header_map(
            target_ws,
            block_start,
            header_row
        )

        date_col_idx = column_index_from_string(
            str(block_start)
        )

        shift_col_idx = header_map.get(
            "shift"
        )

        # ----------------------------------------------------
        # Find target Date + Shift row
        # ----------------------------------------------------

        target_row, seen = find_date_shift_row(
            target_ws,
            date_col_idx,
            shift_col_idx,
            header_row,
            entry["date"],
            entry.get("shift")
        )

        if target_row is None:

            result["status"] = "SKIP"

            result["reason"] = (
                f"no matching Date+Shift "
                f"row in '{sheet_found}'"
            )

            results.append(
                result
            )

            continue

        # ----------------------------------------------------
        # Write fields
        # ----------------------------------------------------

        fields_written = []

        fields_skipped = []

        fields_protected = []

        for field, value in entry.items():

            if field in (
                "row",
                "date",
                "uid",
                "shift",
                "source_file",
                "block_start_col"
            ):

                continue

            target_col = header_map.get(
                normalize(field)
            )

            # ------------------------------------------------
            # Target column doesn't exist
            # ------------------------------------------------

            if target_col is None:

                fields_skipped.append(
                    field
                )

                continue

            target_cell = target_ws.cell(
                row=target_row,
                column=target_col
            )

            existing = target_cell.value

            # ------------------------------------------------
            # Protect formulas
            # ------------------------------------------------

            if (
                isinstance(
                    existing,
                    str
                )
                and existing.strip().startswith("=")
            ):

                fields_protected.append(
                    (
                        field,
                        get_column_letter(
                            target_col
                        ),
                        target_row,
                        f"formula {existing!r}"
                    )
                )

                continue

            # ------------------------------------------------
            # Protect merged cells
            # ------------------------------------------------

            if isinstance(
                target_cell,
                MergedCell
            ):

                fields_protected.append(
                    (
                        field,
                        get_column_letter(
                            target_col
                        ),
                        target_row,
                        "merged cell "
                        "(not the anchor)"
                    )
                )

                continue

            # ------------------------------------------------
            # Add to write list
            # ------------------------------------------------

            fields_written.append(
                (
                    field,
                    value,
                    get_column_letter(
                        target_col
                    ),
                    target_row
                )
            )

            if not DRY_RUN:

                try:

                    target_cell.value = value

                except AttributeError:

                    fields_written.pop()

                    fields_protected.append(
                        (
                            field,
                            get_column_letter(
                                target_col
                            ),
                            target_row,
                            "merged cell "
                            "(write rejected)"
                        )
                    )

        # ----------------------------------------------------
        # Result
        # ----------------------------------------------------

        result["status"] = "OK"

        result["sheet"] = sheet_found

        result["target_row"] = target_row

        result["fields_written"] = len(
            fields_written
        )

        result["fields_skipped"] = len(
            fields_skipped
        )

        result[
            "skipped_field_names"
        ] = fields_skipped

        result[
            "fields_protected"
        ] = fields_protected

        results.append(
            result
        )

        # ----------------------------------------------------
        # Protected fields
        # ----------------------------------------------------

        for (
            field,
            col_letter,
            row_num,
            reason
        ) in fields_protected:

            print(
                f"    {Fore.YELLOW}"
                f"[protected] "
                f"{uid}: {field} "
                f"at {col_letter}{row_num} "
                f"- {reason} "
                f"- left untouched"
                f"{Style.RESET_ALL}"
            )

        # ----------------------------------------------------
        # Progress
        # ----------------------------------------------------

        tick = (
            "+"
            if DRY_RUN
            else "*"
        )

        extra = ""

        if fields_skipped:

            extra = (
                ", "
                + str(
                    len(fields_skipped)
                )
                + " n/a"
            )

        print(
            f"  {Fore.GREEN}"
            f"[{tick}] "
            f"{str(uid):<32} "
            f"-> "
            f"{os.path.basename(filename):<22} "
            f"'{sheet_found}' "
            f"row {target_row} "
            f"({len(fields_written)} "
            f"fields{extra})"
            f"{Style.RESET_ALL}"
        )

        # ----------------------------------------------------
        # Mark master Pushed = Y
        # ----------------------------------------------------

        if src_ws_write is not None:

            src_ws_write.cell(
                row=entry["row"],
                column=pushed_col,
                value="Y"
            )

    # ========================================================
    # PRINT SKIPPED
    # ========================================================

    for r in results:

        if r["status"] == "SKIP":

            print(
                f"  {Fore.RED}"
                f"[x] "
                f"{str(r['uid']):<32} "
                f"-> "
                f"{str(r['file']):<22} "
                f"SKIPPED: "
                f"{r['reason']}"
                f"{Style.RESET_ALL}"
            )

    # ========================================================
    # SAVE SOURCE WORKBOOKS
    # ========================================================

    if not DRY_RUN:

        print()

        for filepath, target_wb in open_targets.items():

            if target_wb is None:
                continue

            log(
                f"Saving "
                f"{os.path.basename(filepath)} ..."
            )

            save_with_retry(
                target_wb,
                filepath
            )

            print(
                f"  Saved: {filepath}"
            )

        # ----------------------------------------------------
        # Save master
        # ----------------------------------------------------

        if master_wb_write is not None:

            log(
                "Saving the master ..."
            )

            save_with_retry(
                master_wb_write,
                MASTER_FILE
            )

            print(
                f"  Saved: {MASTER_FILE}"
            )

    # ========================================================
    # CLOSE WORKBOOKS
    # ========================================================

    for target_wb in open_targets.values():

        if target_wb is None:
            continue

        try:

            target_wb.close()

        except Exception:

            pass

    if master_wb_write is not None:

        try:

            master_wb_write.close()

        except Exception:

            pass

    # ========================================================
    # REPORT
    # ========================================================

    print_report(
        results
    )

    if DRY_RUN:

        print()

        print(
            "This was a DRY RUN."
        )

        print(
            "Review the report above."
        )

        print(
            "Then set "
            "'push_tab_b.dry_run' "
            "to false in config.json "
            "to actually write."
        )

    return target_date


# ============================================================
# REPORT
# ============================================================

def print_report(
    results
):

    ok = [
        r
        for r in results
        if r["status"] == "OK"
    ]

    skip = [
        r
        for r in results
        if r["status"] == "SKIP"
    ]

    print()
    print(
        "=" * 64
    )

    print(
        "  REPORT"
    )

    print(
        "=" * 64
    )

    print(
        f"  Total rows:        "
        f"{len(results)}"
    )

    print(
        f"  Updated:           "
        f"{len(ok)}"
    )

    print(
        f"  Skipped:           "
        f"{len(skip)}"
    )

    # ========================================================
    # BY FILE
    # ========================================================

    files = {}

    for r in results:

        files.setdefault(
            r["file"],
            []
        ).append(r)

    print()
    print(
        "  BY FILE"
    )

    print(
        "  " + "-" * 60
    )

    for filename, rows_for_file in sorted(
        files.items()
    ):

        ok_count = sum(
            1
            for r in rows_for_file
            if r["status"] == "OK"
        )

        skip_count = sum(
            1
            for r in rows_for_file
            if r["status"] == "SKIP"
        )

        if skip_count == 0:

            flag = "OK"
            color = Fore.GREEN

        elif ok_count == 0:

            flag = "NOT UPDATED"
            color = Fore.RED

        else:

            flag = "PARTIAL"
            color = Fore.YELLOW

        print(
            f"  {color}"
            f"[{flag:<11}] "
            f"{str(filename):<24} "
            f"{ok_count} updated, "
            f"{skip_count} skipped"
            f"{Style.RESET_ALL}"
        )

    # ========================================================
    # FILES NOT UPDATED
    # ========================================================

    not_updated = [

        f

        for f, rows_for_file
        in files.items()

        if all(
            r["status"] == "SKIP"
            for r in rows_for_file
        )
    ]

    if not_updated:

        print()

        print(
            f"  {Fore.RED}"
            f"FILES NOT UPDATED AT ALL "
            f"- needs attention:"
            f"{Style.RESET_ALL}"
        )

        print(
            "  " + "-" * 60
        )

        for filename in sorted(
            not_updated
        ):

            reasons = {
                r["reason"]
                for r in files[
                    filename
                ]
                if r["status"] == "SKIP"
            }

            for reason in reasons:

                print(
                    f"    {Fore.RED}"
                    f"- {filename}: "
                    f"{reason}"
                    f"{Style.RESET_ALL}"
                )

    else:

        print()

        print(
            f"  {Fore.GREEN}"
            f"All files had at least "
            f"one row updated."
            f"{Style.RESET_ALL}"
        )

    # ========================================================
    # MISSING FIELDS
    # ========================================================

    fields_missing = {}

    for r in ok:

        for fname in r.get(
            "skipped_field_names",
            []
        ):

            fields_missing.setdefault(
                fname,
                []
            ).append(
                r["uid"]
            )

    if fields_missing:

        print()

        print(
            f"  {Fore.YELLOW}"
            f"FIELDS WITH NO MATCHING "
            f"COLUMN:"
            f"{Style.RESET_ALL}"
        )

        print(
            "  " + "-" * 60
        )

        for fname, uids in sorted(
            fields_missing.items()
        ):

            print(
                f"    {Fore.YELLOW}"
                f"- {fname}: "
                f"{len(uids)} row(s)"
                f"{Style.RESET_ALL}"
            )

    print(
        "=" * 64
    )


# ============================================================
# PROGRAM START
# ============================================================

if __name__ == "__main__":

    last_date = None

    while True:

        last_date = main(
            remembered_date=last_date
        )

        print()

        again = input(
            "Rerun for the same date? "
            "(Y/N): "
        ).strip().lower()

        if again != "y":

            print(
                "Exiting."
            )

            break

        print()