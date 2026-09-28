"""
Push Tab A
==========

Push data from a master "log sheet A" tab into the matching source logsheet
workbooks.

Everything that can change between setups lives in config/config.json
(see the "push_tab_a" section).  Every key is optional - if it is missing the
default used by the previous version of this script applies.

Flow:
    Master sheet row
        | Date + Source File Name + UID + Block Start Column (+ Shift)
        v
    Source workbook -> month sheet -> equipment block -> date/shift row
        v
    Write the configured fields, then mark the master row as pushed

Config keys (all under "push_tab_a", paths under "paths"):

    paths.master_file / paths.folder_a

    dry_run                  true = preview only, nothing is saved
    master_sheet             sheet name inside the master workbook
    master_header_row        row that holds the master column headers
    max_header_search_rows   rows searched for a block's "Date" header
    max_date_search_rows     rows searched for the target date/shift row
    max_block_width          columns scanned to the right of block start
    max_left_scan            columns scanned to the left of block start
    max_header_gap           empty header cells tolerated inside a block
    max_empty_streak         stop reading the master after this many
                             consecutive rows without a valid date

    master_columns           master header names:
                             date, source_file, uid, block_start, shift, pushed
    block_headers            header names inside a block: date, shift
    fields                   list of master columns to push (same header
                             name is looked up inside the block)
    pushed_flag              value written to the "pushed" column (default "Y")
    skip_empty_values        true = never overwrite a target cell with a blank
    file_extensions          workbook types indexed in the source folder
    sheet_name_formats       strftime patterns used to find the month sheet
    date_formats             strptime patterns for dates found in cells
    input_date_formats       strptime patterns accepted at the date prompt
    date                     null = ask at start, "today", or "DD/MM/YYYY"
    repeat_prompt            true = ask "run again for the same date?"
"""

from pathlib import Path
from datetime import datetime, date
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.cell.cell import MergedCell

import io
import json
import os
import re
import time


# ============================================================
# PROJECT / CONFIGURATION
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

# Optional override: set the PUSH_CONFIG_FILE environment variable to use a
# different config file (handy for testing or several setups).
CONFIG_FILE = Path(
    os.environ.get("PUSH_CONFIG_FILE")
    or PROJECT_ROOT / "config" / "config.json"
)


def load_config():
    """Load central project configuration."""

    if not CONFIG_FILE.exists():
        print()
        print("=" * 70)
        print("ERROR: config.json was not found.")
        print()
        print("Expected location:")
        print(CONFIG_FILE)
        print("=" * 70)
        raise SystemExit(1)

    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            return json.load(f)

    except json.JSONDecodeError as e:
        print()
        print("=" * 70)
        print("ERROR: config.json contains invalid JSON.")
        print()
        print(e)
        print("=" * 70)
        raise SystemExit(1)

    except Exception as e:
        print()
        print("=" * 70)
        print("ERROR: Could not read config.json.")
        print()
        print(e)
        print("=" * 70)
        raise SystemExit(1)


CONFIG = load_config()

PATH_CONFIG = CONFIG.get("paths", {})
PUSH_CONFIG = CONFIG.get("push_tab_a", {})


# ============================================================
# CONFIG HELPERS
# ============================================================

def resolve_config_path(value, default_relative):
    """
    Resolve a path from config.

    Absolute path  -> used exactly as supplied.
    Relative path  -> resolved relative to PROJECT_ROOT.
    Missing/empty  -> default_relative is used.
    """

    if value is None or str(value).strip() == "":
        value = default_relative

    path = Path(str(value))

    if path.is_absolute():
        return path

    return PROJECT_ROOT / path


def config_int(value, default):
    """Safely convert a configuration value to int."""

    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def config_bool(value, default=False):
    """Safely convert a configuration value to bool (true/yes/1/on...)."""

    if isinstance(value, bool):
        return value

    if value is None:
        return default

    if isinstance(value, str):
        value = value.strip().lower()

        if value in {"1", "true", "yes", "y", "on"}:
            return True

        if value in {"0", "false", "no", "n", "off"}:
            return False

    return default


def config_list(value, default):
    """Return a non-empty list from config, else the default."""

    if isinstance(value, list) and value:
        return list(value)

    return list(default)


def config_str(value, default):
    """Return a non-empty string from config, else the default."""

    if value is None or str(value).strip() == "":
        return default

    return str(value).strip()


# ============================================================
# DEFAULTS (used when a key is missing from config.json)
# ============================================================

DEFAULT_FIELDS = [
    "Starting HMR",
    "Closing HMR",
    "Starting KMR",
    "Closing KMR",
    "Breakdown Start Time",
    "Breakdown End Time",
    "Breakdown Hours",
    "Diesel Received (Ltr)",
    "Diesel Issued (Ltr)",
    "HMR at Diesel Fill",
    "Hours Run Since Last Fill",
    "Avg. Diesel Consumption (Ltr/Hr)",
    "Container Start Count",
    "Container End Count",
    "Remark",
]

DEFAULT_MASTER_COLUMNS = {
    "date": "Date",
    "source_file": "Source File Name",
    "uid": "UID",
    "block_start": "Block Start Column",
    "shift": "Shift",
    "pushed": "Pushed",
}

DEFAULT_BLOCK_HEADERS = {
    "date": "Date",
    "shift": "Shift",
}

DEFAULT_DATE_FORMATS = [
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%Y-%m-%d",
    "%m/%d/%Y",
    "%d/%m/%y",
    "%d-%m-%y",
]

DEFAULT_INPUT_DATE_FORMATS = [
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d/%m/%y",
    "%d-%m-%y",
    "%Y-%m-%d",
]

DEFAULT_SHEET_NAME_FORMATS = [
    "%b-%y",
    "%B-%y",
    "%b-%Y",
    "%B-%Y",
    "%b%y",
]


# ============================================================
# PATHS FROM CONFIG
# ============================================================

MASTER_FILE = resolve_config_path(
    PATH_CONFIG.get("master_file"),
    "data/Master_Sheet2.xlsm"
)

FOLDER_A = resolve_config_path(
    PATH_CONFIG.get("folder_a"),
    "data/Logsheet-All Site2"
)


# ============================================================
# PUSH TAB A SETTINGS
# ============================================================

DRY_RUN = config_bool(PUSH_CONFIG.get("dry_run"), False)

MASTER_SHEET_A = PUSH_CONFIG.get("master_sheet", "MASTER_LOG_SHEET_A")
MASTER_HEADER_ROW = config_int(PUSH_CONFIG.get("master_header_row"), 1)

MAX_HEADER_SEARCH_ROWS = config_int(PUSH_CONFIG.get("max_header_search_rows"), 25)
MAX_DATE_SEARCH_ROWS = config_int(PUSH_CONFIG.get("max_date_search_rows"), 70)
MAX_BLOCK_WIDTH = config_int(PUSH_CONFIG.get("max_block_width"), 20)
MAX_LEFT_SCAN = config_int(PUSH_CONFIG.get("max_left_scan"), 6)
MAX_HEADER_GAP = config_int(PUSH_CONFIG.get("max_header_gap"), 0)
MAX_EMPTY_STREAK = config_int(PUSH_CONFIG.get("max_empty_streak"), 300)

MASTER_COLUMNS = {
    **DEFAULT_MASTER_COLUMNS,
    **(PUSH_CONFIG.get("master_columns") or {}),
}

BLOCK_HEADERS = {
    **DEFAULT_BLOCK_HEADERS,
    **(PUSH_CONFIG.get("block_headers") or {}),
}

FIELD_NAMES = config_list(PUSH_CONFIG.get("fields"), DEFAULT_FIELDS)

PUSHED_FLAG = config_str(PUSH_CONFIG.get("pushed_flag"), "Y")
SKIP_EMPTY_VALUES = config_bool(PUSH_CONFIG.get("skip_empty_values"), False)

FILE_EXTENSIONS = tuple(
    ext.lower() if ext.startswith(".") else "." + ext.lower()
    for ext in config_list(
        PUSH_CONFIG.get("file_extensions"),
        [".xlsx", ".xlsm"]
    )
)

SHEET_NAME_FORMATS = config_list(
    PUSH_CONFIG.get("sheet_name_formats"),
    DEFAULT_SHEET_NAME_FORMATS
)

DATE_FORMATS = config_list(
    PUSH_CONFIG.get("date_formats"),
    DEFAULT_DATE_FORMATS
)

INPUT_DATE_FORMATS = config_list(
    PUSH_CONFIG.get("input_date_formats"),
    DEFAULT_INPUT_DATE_FORMATS
)

RUN_DATE_SETTING = PUSH_CONFIG.get("date")
REPEAT_PROMPT = config_bool(PUSH_CONFIG.get("repeat_prompt"), True)


# ============================================================
# COLORAMA (optional)
# ============================================================

try:
    from colorama import init, Fore, Style

    init(autoreset=True)

    GREEN = Fore.GREEN
    YELLOW = Fore.YELLOW
    RED = Fore.RED
    CYAN = Fore.CYAN
    MAGENTA = Fore.MAGENTA
    RESET = Style.RESET_ALL

except Exception:

    GREEN = ""
    YELLOW = ""
    RED = ""
    CYAN = ""
    MAGENTA = ""
    RESET = ""


# ============================================================
# TIMER / LOGGING
# ============================================================

START_TIME = time.time()


def log(message=""):
    elapsed = time.time() - START_TIME
    print(f"[{elapsed:8.1f}s] {message}")


# ============================================================
# BASIC HELPERS
# ============================================================

def read_file_bytes(path):
    """
    Read a complete file into memory.

    Faster than letting openpyxl read it piece by piece when the file
    lives on a network location.
    """

    with open(path, "rb") as f:
        return f.read()


def normalize(value):
    """Normalize text for case/space-insensitive comparison."""

    if value is None:
        return ""

    return re.sub(r"\s+", " ", str(value).strip().lower())


def is_empty(value):
    """True for None or whitespace-only text."""

    return value is None or (
        isinstance(value, str) and value.strip() == ""
    )


# Normalized config names used for comparisons
KEY_DATE = normalize(BLOCK_HEADERS["date"])
KEY_SHIFT = normalize(BLOCK_HEADERS["shift"])
PUSHED_FLAG_KEY = normalize(PUSHED_FLAG)


# ============================================================
# SAVE WITH RETRY
# ============================================================

def save_with_retry(wb, path):
    """Save a workbook; if it is open/locked, wait and retry."""

    path = Path(path)

    while True:

        try:
            wb.save(path)
            log(f"{GREEN}Saved successfully: {path.name}{RESET}")
            return

        except PermissionError:

            print()
            print(f"{YELLOW}Workbook appears to be open or locked:{RESET}")
            print(path)
            print("Please close the Excel file.")
            input("Press ENTER to retry...")

        except Exception as e:

            print()
            print(f"{RED}Error saving workbook:{RESET}")
            print(path)
            print(e)

            choice = input(
                "Press ENTER to retry or type Q to quit: "
            ).strip().lower()

            if choice == "q":
                raise


# ============================================================
# DATE HANDLING
# ============================================================

def normalize_date(value):
    """Convert a cell value to datetime.date (or None)."""

    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    if not text:
        return None

    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass

    return None


def parse_input_date(value):
    """Parse a date typed by the user or set in config."""

    value = str(value).strip()

    for fmt in INPUT_DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass

    return None


def prompt_for_date():

    while True:

        value = input("\nEnter date to process (DD/MM/YYYY): ").strip()

        if not value:
            print(f"{YELLOW}Please enter a date.{RESET}")
            continue

        parsed = parse_input_date(value)

        if parsed:
            return parsed

        print(f"{RED}Invalid date format.{RESET}")
        print("Example: 18/09/2026")


def configured_run_date():
    """
    Date from config ("date": "today" or "DD/MM/YYYY"), else None.
    """

    if is_empty(RUN_DATE_SETTING):
        return None

    text = str(RUN_DATE_SETTING).strip()

    if text.lower() == "today":
        return date.today()

    parsed = parse_input_date(text)

    if parsed is None:
        print(
            f"{YELLOW}WARNING: config 'date' value '{text}' is not "
            f"a valid date. You will be asked instead.{RESET}"
        )

    return parsed


# ============================================================
# FIND MONTH SHEET
# ============================================================

def find_matching_sheet(wb, target_date):
    """
    Find the sheet that belongs to target_date (e.g. Sep-26, September-2026).
    Patterns come from config "sheet_name_formats".
    """

    if target_date is None:
        return None

    if len(wb.sheetnames) == 1:
        return wb[wb.sheetnames[0]]

    wanted = {
        normalize(target_date.strftime(fmt))
        for fmt in SHEET_NAME_FORMATS
    }

    for sheet_name in wb.sheetnames:
        if normalize(sheet_name) in wanted:
            return wb[sheet_name]

    # Loose match: month name + year anywhere in the sheet name
    month_short = target_date.strftime("%b").lower()
    month_long = target_date.strftime("%B").lower()
    year_short = target_date.strftime("%y")
    year_full = target_date.strftime("%Y")

    for sheet_name in wb.sheetnames:

        n = normalize(sheet_name)

        if (
            (month_short in n or month_long in n)
            and (year_short in n or year_full in n)
        ):
            return wb[sheet_name]

    return None


# ============================================================
# LOCATE BLOCK / HEADERS / ROW
# ============================================================

def find_block_header_row(ws, start_col_letter):
    """Find the row that holds the block's 'Date' header."""

    try:
        start_col = column_index_from_string(start_col_letter)
    except Exception:
        return None

    last_row = min(ws.max_row, MAX_HEADER_SEARCH_ROWS)

    for row in range(1, last_row + 1):

        value = ws.cell(row=row, column=start_col).value

        if normalize(value) == KEY_DATE:
            return row

    return None


def find_block_header_map(ws, start_col_letter, header_row):
    """
    Map normalized header name -> column number for one equipment block.

    Scans right up to max_block_width and left up to max_left_scan.
    """

    try:
        start_col = column_index_from_string(start_col_letter)
    except Exception:
        return {}

    header_map = {}

    # ---- scan right ----------------------------------------
    gap = 0

    for offset in range(MAX_BLOCK_WIDTH):

        col = start_col + offset

        if col > ws.max_column:
            break

        value = ws.cell(row=header_row, column=col).value

        if is_empty(value):
            gap += 1

            if gap > MAX_HEADER_GAP:
                break

            continue

        gap = 0

        key = normalize(value)

        if key:
            header_map.setdefault(key, col)

    # ---- scan left -----------------------------------------
    gap = 0

    for offset in range(1, MAX_LEFT_SCAN + 1):

        col = start_col - offset

        if col < 1:
            break

        value = ws.cell(row=header_row, column=col).value

        if is_empty(value):
            gap += 1

            if gap > MAX_HEADER_GAP:
                break

            continue

        gap = 0

        key = normalize(value)

        if key:
            header_map.setdefault(key, col)

    return header_map


def find_date_shift_row(ws, date_col, shift_col, target_date, target_shift):
    """Find the row that matches date (+ shift when the block has one)."""

    if target_date is None:
        return None

    wanted_shift = normalize(target_shift)

    last_row = min(ws.max_row, MAX_DATE_SEARCH_ROWS)

    for row in range(1, last_row + 1):

        cell_date = normalize_date(
            ws.cell(row=row, column=date_col).value
        )

        if cell_date != target_date:
            continue

        if shift_col is None:
            return row

        shift_value = ws.cell(row=row, column=shift_col).value

        if normalize(shift_value) == wanted_shift:
            return row

    return None


def locate_target(wb, target_date, block_start_column, shift):
    """
    Walk: month sheet -> block header row -> header map -> date/shift row.

    Returns (ws, header_map, target_row, error_reason).
    error_reason is None on success.
    """

    ws = find_matching_sheet(wb, target_date)

    if ws is None:
        return None, None, None, "Month sheet not found"

    header_row = find_block_header_row(ws, block_start_column)

    if header_row is None:
        return ws, None, None, "Block header not found"

    header_map = find_block_header_map(ws, block_start_column, header_row)

    if not header_map:
        return ws, None, None, "No block headers"

    date_col = header_map.get(KEY_DATE)

    if date_col is None:
        return ws, header_map, None, "Date column not found"

    shift_col = header_map.get(KEY_SHIFT)

    target_row = find_date_shift_row(
        ws, date_col, shift_col, target_date, shift
    )

    if target_row is None:
        return ws, header_map, None, "Date + shift row not found"

    return ws, header_map, target_row, None


# ============================================================
# SOURCE FILE INDEX
# ============================================================

def build_file_index(root_folder):
    """
    Recursively index workbooks in root_folder.

    Returns {"filename_without_extension": [Path, Path, ...]}
    """

    root_folder = Path(root_folder)
    index = {}

    if not root_folder.exists():
        return index

    log("Building source file index from:")
    log(f"  {root_folder}")

    for root, dirs, files in os.walk(root_folder):

        for filename in files:

            if filename.startswith("~$"):
                continue

            if not filename.lower().endswith(FILE_EXTENSIONS):
                continue

            stem = Path(filename).stem.lower()

            index.setdefault(stem, []).append(Path(root) / filename)

    # Shallow folders first
    for key in index:
        index[key].sort(key=lambda p: (len(p.parts), str(p).lower()))

    log(f"Indexed {sum(len(v) for v in index.values())} Excel files.")

    return index


def find_source_file(root_folder, source_name, index):
    """
    Find a source workbook.

    Handles relative paths, names with/without extension, any letter case.
    """

    if not source_name:
        return None

    root_folder = Path(root_folder)
    source_text = str(source_name).strip()

    if not source_text:
        return None

    # ---- direct relative path ------------------------------
    candidate = Path(source_text)

    if not candidate.is_absolute():

        candidate = root_folder / candidate

        if candidate.exists() and candidate.is_file():
            return candidate

        for ext in FILE_EXTENSIONS:

            with_ext = Path(str(candidate) + ext)

            if with_ext.exists() and with_ext.is_file():
                return with_ext

    # ---- name matching -------------------------------------
    source_stem = Path(source_text).stem.lower()

    matches = list(index.get(source_stem, []))

    if not matches:
        return None

    if len(matches) > 1:

        print()
        print(
            f"{YELLOW}WARNING: Multiple source files found for "
            f"'{source_name}'{RESET}"
        )

        for p in matches:
            print(f"   {p}")

        print(f"{YELLOW}Using:{RESET}")
        print(f"   {matches[0]}")

    return matches[0]


# ============================================================
# READ MASTER ROWS  (streaming - fast on large sheets)
# ============================================================

def read_master_rows():
    """
    Read pending rows from the master sheet.

    Streams the sheet once with iter_rows instead of calling ws.cell() per
    cell, and stops after max_empty_streak consecutive rows without a valid
    date.  Rows already marked as pushed are skipped.
    """

    log("Reading master workbook:")
    log(f"  {MASTER_FILE}")

    if not MASTER_FILE.exists():
        print()
        print(f"{RED}ERROR: Master file not found:{RESET}")
        print(MASTER_FILE)
        raise SystemExit(1)

    try:
        wb = load_workbook(
            filename=io.BytesIO(read_file_bytes(MASTER_FILE)),
            read_only=True,
            data_only=True
        )

    except Exception as e:
        print()
        print(f"{RED}ERROR: Could not open master workbook:{RESET}")
        print(e)
        raise SystemExit(1)

    log("Workbook opened, reading headers...")

    if MASTER_SHEET_A not in wb.sheetnames:
        print()
        print(
            f"{RED}ERROR: Sheet '{MASTER_SHEET_A}' not found in "
            f"master workbook.{RESET}"
        )
        print()
        print("Available sheets:")

        for name in wb.sheetnames:
            print(f"   {name}")

        wb.close()
        raise SystemExit(1)

    ws = wb[MASTER_SHEET_A]

    # ---- headers -------------------------------------------
    header_values = next(
        ws.iter_rows(
            min_row=MASTER_HEADER_ROW,
            max_row=MASTER_HEADER_ROW,
            values_only=True
        ),
        ()
    )

    headers = {}

    for col, value in enumerate(header_values, start=1):
        key = normalize(value)
        if key:
            headers.setdefault(key, col)

    cols = {
        name: headers.get(normalize(header_text))
        for name, header_text in MASTER_COLUMNS.items()
    }

    required = ["date", "source_file", "uid", "block_start"]

    missing = [MASTER_COLUMNS[name] for name in required if not cols.get(name)]

    if missing:
        print()
        print(
            f"{RED}ERROR: Required columns missing from master "
            f"sheet:{RESET}"
        )

        for header_text in missing:
            print(f"   {header_text}")

        wb.close()
        raise SystemExit(1)

    field_columns = {}

    for field in FIELD_NAMES:
        col = headers.get(normalize(field))
        if col:
            field_columns[field] = col

    not_in_master = [f for f in FIELD_NAMES if f not in field_columns]

    if not_in_master:
        log(
            f"{YELLOW}Fields not found in master header row "
            f"(will not be pushed): {', '.join(not_in_master)}{RESET}"
        )

    # ---- rows ----------------------------------------------
    needed = [c for c in cols.values() if c] + list(field_columns.values())
    last_col = max(needed)
    first_data_row = MASTER_HEADER_ROW + 1

    def cell(row, col):
        if not col or col > len(row):
            return None
        return row[col - 1]

    master_rows = []
    empty_streak = 0

    for row_number, row in enumerate(
        ws.iter_rows(
            min_row=first_data_row,
            max_col=last_col,
            values_only=True
        ),
        start=first_data_row
    ):

        row_date = normalize_date(cell(row, cols["date"]))

        if row_date is None:
            empty_streak += 1

            if empty_streak >= MAX_EMPTY_STREAK:
                log(
                    f"Stopped reading at row {row_number}: "
                    f"{MAX_EMPTY_STREAK} rows in a row without a date."
                )
                break

            continue

        empty_streak = 0

        source_file_name = cell(row, cols["source_file"])
        uid = cell(row, cols["uid"])
        block_start_column = cell(row, cols["block_start"])

        if (
            is_empty(source_file_name)
            or is_empty(uid)
            or is_empty(block_start_column)
        ):
            continue

        pushed_value = cell(row, cols.get("pushed"))

        if cols.get("pushed") and normalize(pushed_value) == PUSHED_FLAG_KEY:
            continue

        shift_value = cell(row, cols.get("shift")) if cols.get("shift") else ""

        master_rows.append(
            {
                "excel_row": row_number,
                "date": row_date,
                "source_file_name": str(source_file_name).strip(),
                "uid": str(uid).strip(),
                "shift": str(
                    shift_value if shift_value is not None else ""
                ).strip(),
                "block_start_column": str(block_start_column).strip(),
                "fields": {
                    name: cell(row, col)
                    for name, col in field_columns.items()
                },
                "pushed": pushed_value,
            }
        )

    wb.close()

    log(f"Read {len(master_rows)} pending master rows.")

    return master_rows


# ============================================================
# MAIN PROCESS
# ============================================================

def print_banner(title, subtitle=""):

    print()
    print("=" * 80)
    print(f" {title}")

    if subtitle:
        print(f" {subtitle}")

    print("=" * 80)


def main(remembered_date=None):

    global START_TIME

    START_TIME = time.time()

    print_banner("PUSH TAB A", f"{MASTER_SHEET_A}  -->  SOURCE LOGSHEETS")

    print()
    print("Project root:")
    print(f"  {PROJECT_ROOT}")

    print()
    print("Configuration:")
    print(f"  {CONFIG_FILE}")

    print()
    print("Master file:")
    print(f"  {MASTER_FILE}")

    print()
    print("Source folder:")
    print(f"  {FOLDER_A}")

    print()
    print(
        "Mode:",
        f"{YELLOW}DRY RUN{RESET}" if DRY_RUN else f"{GREEN}LIVE PUSH{RESET}"
    )

    # ---- validate paths ------------------------------------
    if not MASTER_FILE.exists():
        print()
        print(f"{RED}ERROR: Master file does not exist:{RESET}")
        print(MASTER_FILE)
        return False, None

    if not FOLDER_A.exists() or not FOLDER_A.is_dir():
        print()
        print(f"{RED}ERROR: Source folder is missing or not a folder:{RESET}")
        print(FOLDER_A)
        return False, None

    # ---- read master ---------------------------------------
    master_rows = read_master_rows()

    if not master_rows:
        print()
        print(f"{GREEN}No pending rows found.{RESET}")
        return True, remembered_date

    # ---- date ----------------------------------------------
    if remembered_date is not None:

        target_date = remembered_date

        print()
        print(f"Using remembered date: {target_date.strftime('%d/%m/%Y')}")

    else:

        target_date = configured_run_date() or prompt_for_date()

        print()
        print(f"Processing date: {target_date.strftime('%d/%m/%Y')}")

    rows_for_date = [r for r in master_rows if r["date"] == target_date]

    if not rows_for_date:
        print()
        print(
            f"{YELLOW}No pending rows found for "
            f"{target_date.strftime('%d/%m/%Y')}.{RESET}"
        )
        return True, target_date

    print()
    print(f"Rows to process: {len(rows_for_date)}")

    # ---- source file index ---------------------------------
    file_index = build_file_index(FOLDER_A)

    if not file_index:
        print()
        print(f"{RED}ERROR: No Excel files found in source folder.{RESET}")
        return False, target_date

    # ---- open master for the Pushed update -----------------
    master_wb = None
    master_ws = None
    pushed_column = None

    if not DRY_RUN:

        try:
            log("Opening master workbook for update (can take a while)...")

            master_wb = load_workbook(
                MASTER_FILE,
                keep_vba=str(MASTER_FILE).lower().endswith(".xlsm")
            )

            master_ws = master_wb[MASTER_SHEET_A]

            wanted = normalize(MASTER_COLUMNS["pushed"])

            for col in range(1, master_ws.max_column + 1):

                value = master_ws.cell(
                    row=MASTER_HEADER_ROW,
                    column=col
                ).value

                if normalize(value) == wanted:
                    pushed_column = col
                    break

            if pushed_column is None:
                print()
                print(
                    f"{YELLOW}WARNING: '{MASTER_COLUMNS['pushed']}' column "
                    f"not found in master.{RESET}"
                )
                print(
                    "Rows will still be pushed, but their status cannot "
                    "be updated."
                )

            log("Master workbook ready.")

        except Exception as e:
            print()
            print(f"{RED}ERROR: Could not open master for writing:{RESET}")
            print(e)
            return False, target_date

    # ---- process rows --------------------------------------
    # cache key -> {"wb": workbook, "path": Path, "changed": bool}
    workbook_cache = {}

    updated_rows = []
    skipped_rows = []
    missing_fields = {}

    def skip(row_number, reason, source):
        skipped_rows.append(
            {"row": row_number, "reason": reason, "source": source}
        )

    for index, row_data in enumerate(rows_for_date, start=1):

        master_row_number = row_data["excel_row"]
        source_file_name = row_data["source_file_name"]
        shift = row_data["shift"]
        block_start_column = row_data["block_start_column"]

        print()
        print("-" * 80)
        print(f"[{index}/{len(rows_for_date)}] Master Row {master_row_number}")
        print(f"  UID: {row_data['uid']}")
        print(f"  Source: {source_file_name}")
        print(f"  Shift: {shift}")
        print(f"  Block Start: {block_start_column}")

        # ---- source workbook -------------------------------
        source_path = find_source_file(FOLDER_A, source_file_name, file_index)

        if source_path is None:
            print(f"{RED}  SOURCE FILE NOT FOUND{RESET}")
            skip(master_row_number, "Source file not found", source_file_name)
            continue

        print("  Source path:")
        print(f"    {source_path}")

        cache_key = str(source_path).lower()

        if cache_key not in workbook_cache:

            try:
                log(f"Opening target workbook: {source_path.name}")

                workbook_cache[cache_key] = {
                    "wb": load_workbook(
                        source_path,
                        keep_vba=source_path.suffix.lower() == ".xlsm"
                    ),
                    "path": source_path,
                    "changed": False,
                }

            except Exception as e:
                print(f"{RED}  ERROR opening workbook:{RESET}")
                print(e)
                skip(master_row_number, "Workbook open error", source_file_name)
                continue

        target_wb = workbook_cache[cache_key]["wb"]

        # ---- month sheet / block / row ---------------------
        target_ws, header_map, target_row, error = locate_target(
            target_wb, target_date, block_start_column, shift
        )

        if error:
            print(f"{RED}  {error.upper()}{RESET}")

            if error == "Month sheet not found":
                print(f"  Target date: {target_date}")
                print("  Available sheets:")
                for name in target_wb.sheetnames:
                    print(f"    {name}")

            elif error == "Date + shift row not found":
                print(f"  Date: {target_date}")
                print(f"  Shift: {shift}")

            skip(master_row_number, error, source_file_name)
            continue

        print(f"  Sheet: {target_ws.title}")
        print(f"  Target row: {target_row}")

        # ---- write fields ----------------------------------
        row_changed = False
        row_missing_fields = []

        for field_name, value in row_data["fields"].items():

            target_col = header_map.get(normalize(field_name))

            if target_col is None:
                row_missing_fields.append(field_name)
                continue

            if SKIP_EMPTY_VALUES and is_empty(value):
                continue

            cell = target_ws.cell(row=target_row, column=target_col)

            if isinstance(cell, MergedCell):
                print(f"  SKIP merged cell: {field_name}")
                continue

            if isinstance(cell.value, str) and cell.value.startswith("="):
                print(f"  SKIP formula: {field_name}")
                continue

            print(f"  {field_name}: {cell.coordinate} = {value!r}")

            if not DRY_RUN:
                cell.value = value

            row_changed = True

        if row_missing_fields:

            missing_fields[source_file_name] = sorted(
                set(missing_fields.get(source_file_name, []) + row_missing_fields)
            )

            print()
            print(f"{YELLOW}  Missing target fields:{RESET}")

            for field in row_missing_fields:
                print(f"    - {field}")

        # ---- result ----------------------------------------
        if row_changed:

            workbook_cache[cache_key]["changed"] = True

            if (
                not DRY_RUN
                and master_ws is not None
                and pushed_column is not None
            ):
                master_ws.cell(
                    row=master_row_number,
                    column=pushed_column
                ).value = PUSHED_FLAG

            updated_rows.append(
                {
                    "row": master_row_number,
                    "source": source_file_name,
                    "target_row": target_row,
                    "sheet": target_ws.title,
                }
            )

        else:
            skip(master_row_number, "No fields updated", source_file_name)

    # ---- save source workbooks -----------------------------
    print_banner("SAVING SOURCE WORKBOOKS")

    if DRY_RUN:

        print()
        print(f"{YELLOW}DRY RUN enabled - no workbooks will be saved.{RESET}")

    else:

        for entry in workbook_cache.values():

            if not entry["changed"]:
                continue

            print()
            print("Saving:")
            print(f"  {entry['path']}")

            save_with_retry(entry["wb"], entry["path"])

    # ---- save master ---------------------------------------
    if not DRY_RUN and master_wb is not None:

        print()
        print("Saving master workbook...")

        save_with_retry(master_wb, MASTER_FILE)

    # ---- close ---------------------------------------------
    for entry in workbook_cache.values():
        try:
            entry["wb"].close()
        except Exception:
            pass

    if master_wb is not None:
        try:
            master_wb.close()
        except Exception:
            pass

    print_report(target_date, updated_rows, skipped_rows, missing_fields)

    return True, target_date


# ============================================================
# REPORT
# ============================================================

def print_report(target_date, updated_rows, skipped_rows, missing_fields):

    print_banner("PUSH REPORT")

    print()
    print("Date:", target_date.strftime("%d/%m/%Y"))
    print("Updated rows:", len(updated_rows))
    print("Skipped rows:", len(skipped_rows))

    if DRY_RUN:
        print()
        print(
            f"{YELLOW}NOTE: DRY RUN was enabled. "
            f"No Excel files were changed.{RESET}"
        )

    if updated_rows:
        print()
        print(f"{GREEN}UPDATED ROWS{RESET}")

        for item in updated_rows:
            print(
                f"  Master Row {item['row']} -> {item['source']} "
                f"-> {item['sheet']} -> Row {item['target_row']}"
            )

    if skipped_rows:
        print()
        print(f"{YELLOW}SKIPPED ROWS{RESET}")

        for item in skipped_rows:
            print(
                f"  Master Row {item['row']} -> {item['reason']} "
                f"-> {item['source']}"
            )

    if missing_fields:
        print()
        print(f"{YELLOW}MISSING TARGET FIELDS{RESET}")

        for source, fields in missing_fields.items():
            print()
            print(f"  {source}")

            for field in fields:
                print(f"     - {field}")

    print_banner("DONE")
    print()


# ============================================================
# PROGRAM START
# ============================================================

if __name__ == "__main__":

    remembered_date = None

    while True:

        success, processed_date = main(remembered_date)

        if processed_date is not None:
            remembered_date = processed_date

        if not REPEAT_PROMPT:
            break

        print()

        answer = input(
            "Run Push A again for the SAME date? (Y/N): "
        ).strip().lower()

        if answer != "y":
            break

    print()
    print("Push A closed.")